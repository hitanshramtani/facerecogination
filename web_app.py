from datetime import datetime
from pathlib import Path
import subprocess
import sys

import cv2
import numpy as np
import streamlit as st
from openpyxl import load_workbook

from register_face import register_face_from_image


APP_FILE = Path("app.py")
ATTENDANCE_FILE = Path("Attendance_sheet_demo1.xlsx")
NAMES_FILE = Path("student_names2.npy")
PRESENT_VALUE = "P"


def _load_registered_count():
    if not NAMES_FILE.exists():
        return 0
    names = np.load(NAMES_FILE, allow_pickle=True)
    return int(len(names))


def _read_attendance_stats():
    if not ATTENDANCE_FILE.exists():
        return {
            "present_today": 0,
            "absent_today": 0,
            "total_students": 0,
            "daily_labels": [],
            "daily_present": [],
        }

    workbook = load_workbook(ATTENDANCE_FILE)
    sheet = workbook.active
    headers = [sheet.cell(row=1, column=col).value for col in range(2, sheet.max_column + 1)]

    rows = []
    for row_idx in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row_idx, column=1).value
        if not name:
            continue
        row_values = [sheet.cell(row=row_idx, column=col).value for col in range(2, sheet.max_column + 1)]
        rows.append((str(name), row_values))

    today_header = datetime.now().strftime("%d-%b-%Y")
    present_today = 0
    if today_header in headers:
        idx = headers.index(today_header)
        for _, values in rows:
            if str(values[idx]).strip().upper() == PRESENT_VALUE:
                present_today += 1

    total_students = len(rows)
    absent_today = max(total_students - present_today, 0)

    daily_present = []
    daily_labels = []
    for col_idx, header in enumerate(headers):
        if not header:
            continue
        h = str(header)
        if "-Apr-" not in h:
            continue
        count_p = 0
        for _, values in rows:
            value = values[col_idx]
            if str(value).strip().upper() == PRESENT_VALUE:
                count_p += 1
        daily_labels.append(h)
        daily_present.append(count_p)

    workbook.close()
    return {
        "present_today": present_today,
        "absent_today": absent_today,
        "total_students": total_students,
        "daily_labels": daily_labels,
        "daily_present": daily_present,
    }


def _init_state():
    if "page" not in st.session_state:
        st.session_state.page = "home"
    if "attendance_process" not in st.session_state:
        st.session_state.attendance_process = None
    if "register_result" not in st.session_state:
        st.session_state.register_result = None


def _go_home():
    st.session_state.page = "home"


def _go_mark():
    st.session_state.page = "mark"


def _go_register():
    st.session_state.page = "register"


def _is_attendance_running():
    proc = st.session_state.attendance_process
    return proc is not None and proc.poll() is None


def _start_attendance():
    if not APP_FILE.exists():
        st.error("app.py not found.")
        return
    if _is_attendance_running():
        st.info("Attendance is already running.")
        return
    proc = subprocess.Popen([sys.executable, str(APP_FILE)])
    st.session_state.attendance_process = proc
    st.success("Attendance started. Camera window opened from app.py.")


def _stop_attendance():
    proc = st.session_state.attendance_process
    if proc is None:
        st.info("Attendance is not running.")
        return
    if proc.poll() is None:
        proc.terminate()
        try:
            proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            proc.kill()
    st.session_state.attendance_process = None
    st.success("Attendance stopped.")


def render_home():
    st.title("Face Attendance Dashboard")
    st.write("Choose an action from below.")

    stats = _read_attendance_stats()
    registered = _load_registered_count()

    c1, c2, c3 = st.columns(3)
    c1.metric("Registered Students", registered)
    c2.metric("Present Today", stats["present_today"])
    c3.metric("Absent Today", stats["absent_today"])

    st.subheader("Today Snapshot")
    snapshot_data = {
        "Status": ["Present", "Absent"],
        "Count": [stats["present_today"], stats["absent_today"]],
    }
    st.bar_chart(snapshot_data, x="Status", y="Count")

    if stats["daily_labels"]:
        st.subheader("April Daily Present Count")
        chart_data = {"Present": stats["daily_present"]}
        st.line_chart(chart_data)
        st.caption(f"Dates: {', '.join(stats['daily_labels'])}")
    else:
        st.info("No April attendance data found yet.")

    b1, b2 = st.columns(2)
    with b1:
        st.button("Mark Attendance", on_click=_go_mark, use_container_width=True)
    with b2:
        st.button("Register Face", on_click=_go_register, use_container_width=True)


def render_mark_attendance():
    st.title("Mark Attendance")
    st.write("This runs `app.py`. Use Exit button to stop and return to Home.")

    running = _is_attendance_running()
    if running:
        st.success("Attendance session is running.")
    else:
        st.warning("Attendance session is not running.")

    c1, c2 = st.columns(2)
    with c1:
        if st.button("Start Attendance", use_container_width=True):
            _start_attendance()
            st.rerun()
    with c2:
        if st.button("Exit Attendance", use_container_width=True):
            _stop_attendance()
            _go_home()
            st.rerun()

    st.info("If the camera window is active, you can also press 'q' there.")


def render_register_face():
    st.title("Register Face")
    st.write("Enter name, capture photo, then press Register.")

    name = st.text_input("Student Name")
    captured = st.camera_input("Capture student photo")

    if st.button("Register", use_container_width=True):
        if not name.strip():
            st.error("Please enter student name.")
        elif captured is None:
            st.error("Please capture a photo first.")
        else:
            file_bytes = np.asarray(bytearray(captured.read()), dtype=np.uint8)
            image_bgr = cv2.imdecode(file_bytes, cv2.IMREAD_COLOR)
            if image_bgr is None:
                st.error("Failed to decode captured photo.")
            else:
                ok, message = register_face_from_image(image_bgr, name)
                st.session_state.register_result = (ok, message)

    result = st.session_state.register_result
    if result:
        ok, message = result
        if ok:
            st.success(message)
        else:
            st.warning(message)

    if st.button("Back Home", use_container_width=True):
        st.session_state.register_result = None
        _go_home()
        st.rerun()


def main():
    st.set_page_config(page_title="Face Attendance System", page_icon=":camera:", layout="wide")
    _init_state()

    page = st.session_state.page
    if page == "home":
        render_home()
    elif page == "mark":
        render_mark_attendance()
    elif page == "register":
        render_register_face()
    else:
        st.session_state.page = "home"
        render_home()


if __name__ == "__main__":
    main()
